"""Builds Docebo_Reverse_DCF.xlsx - "What is Docebo's stock price secretly betting on?"

A 10-year ARR -> revenue -> free cash flow -> DCF model of Docebo Inc. (TSX: DCBO),
run backwards: it solves for the churn rate that makes the model's value per share
equal today's share price. Everything is a live formula; blue cells are inputs.
Run:  python build_model.py
"""
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FONT = "Arial"
F = lambda color="000000", bold=False, size=10, italic=False: Font(
    name=FONT, size=size, bold=bold, italic=italic, color=color)
F_IN, F_CALC, F_LINK = F("0000FF"), F(), F("008000")
F_HDR = F("FFFFFF", True)
FILL_HDR = PatternFill("solid", fgColor="1F3864")
FILL_SEC = PatternFill("solid", fgColor="D9E1F2")
FILL_KEY = PatternFill("solid", fgColor="FFFF00")
FILL_ANS = PatternFill("solid", fgColor="E2EFDA")
TOP = Border(top=Side(style="thin"))
BOX = Border(*(Side(style="medium"),) * 4)

USDM = '#,##0.0;(#,##0.0);"-"'          # US$ millions, one decimal
USD2 = '$#,##0.00;($#,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'
MULT = '0.0x'
NUM2 = '0.00'

YEARS = 10
YC = [chr(ord("C") + i) for i in range(YEARS)]   # Year 1..10 -> C..L
Y1, Y10 = YC[0], YC[-1]

wb = Workbook()


def put(ws, ref, value, font=F_CALC, fmt=None, fill=None, bold=False, align=None, border=None):
    c = ws[ref]
    c.value = value
    c.font = Font(name=FONT, size=font.size, bold=bold or font.bold, color=font.color,
                  italic=font.italic)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if align:
        c.alignment = Alignment(horizontal=align, wrap_text=False)
    if border:
        c.border = border
    return c


def header(ws, r, labels, start="A"):
    for i, t in enumerate(labels):
        c = ws.cell(r, ord(start) - 64 + i, t)
        c.font, c.fill = F_HDR, FILL_HDR
        c.alignment = Alignment(horizontal="left" if i == 0 else "center", wrap_text=True)


def section(ws, r, text, width):
    ws.cell(r, 1, text).font = F(bold=True)
    for i in range(1, width + 1):
        ws.cell(r, i).fill = FILL_SEC


def note(ws, ref, text):
    c = ws[ref]
    c.value = text
    c.font = F("595959", italic=True, size=9)


def widths(ws, spec):
    for k, v in spec.items():
        ws.column_dimensions[k].width = v


# =============================================================== START HERE
g = wb.active
g.title = "Start Here"
widths(g, {"A": 3, "B": 120})
lines = [
    ("What is Docebo's stock price secretly betting on?", F(bold=True, size=14)),
    ("A reverse DCF at the customer level: which churn rate makes today's share price make sense?", F(italic=True)),
    ("", None),
    ("THE IDEA IN ONE BREATH", F(bold=True)),
    ("Normal model: assumptions -> value.  Reverse model: today's price -> the assumption the market must believe.", None),
    ("We hold everything else at the company's real numbers and ask: how much ARR would have to cancel every year for the stock to be fairly priced?", None),
    ("", None),
    ("HOW THE TABS FLOW", F(bold=True)),
    ("1. Inputs       - real numbers from Docebo's filings (8 quarters) + today's price. Every figure has a source.", None),
    ("2. Assumptions  - the forecast levers. The yellow CHURN cell is the one the reverse DCF solves for.", None),
    ("3. ARR Model    - 10 years: Starting ARR + new ARR + expansion - churn = Ending ARR -> revenue -> EBITDA -> free cash flow.", None),
    ("4. DCF          - shrink each year's cash to today's dollars, add a terminal value, add cash, subtract debt, divide by shares.", None),
    ("5. Reverse DCF  - THE ANSWER: market-implied churn vs. reported churn, solved automatically (no macro needed).", None),
    ("6. Sensitivity  - share price for every combination of churn (down the side) and discount rate (across the top).", None),
    ("7. Tracker      - log each run. After Q3-2026 results (early Nov 2026) update Inputs and add a row: did the bet change?", None),
    ("", None),
    ("COLOR CODE", F(bold=True)),
    ("Blue text = hard-coded input  |  Black = formula  |  Green = link from another tab  |  Yellow = key lever  |  Light green = answer", None),
    ("", None),
    ("DO THE EXCEL SKILLS YOURSELF (interview practice)", F(bold=True)),
    ("GOAL SEEK:  Data -> What-If Analysis -> Goal Seek.  Set cell: DCF!C22   To value: 33.22   By changing: Assumptions!C5.  Click OK.", None),
    ("   Excel changes the churn cell until the model price equals the market price. Compare with the answer on the Reverse DCF tab (they should match).", None),
    ("   Then press Ctrl+Z (or retype 9%) to put churn back.", None),
    ("DATA TABLE:  on an empty sheet, put =DCF!C22 in the top-left cell, churn rates down the column below it, discount rates across the row to its right.", None),
    ("   Select the whole block -> Data -> What-If Analysis -> Data Table.  Row input cell: Assumptions!C13.  Column input cell: Assumptions!C5.", None),
    ("   Your grid should match the Sensitivity tab (which is built with plain formulas so it also works in Google Sheets).", None),
    ("", None),
    ("IMPORTANT CAVEATS", F(bold=True)),
    ("- Figures were collected from Docebo press releases via news/search summaries. Before you share this, re-check each blue number against the filings on sedarplus.ca.", None),
    ("- Docebo reports in US dollars; the TSX price is in Canadian dollars, so the model converts with the USD/CAD rate on Inputs.", None),
    ("- Docebo does not report gross churn. 'Reported churn' = expansion assumption + (1 - reported net revenue retention). If you change expansion, both sides move together.", None),
    ("- Adjusted EBITDA excludes stock-based compensation, which is a real cost to shareholders. That makes this DCF a bit generous.", None),
    ("- Solving for churn only is a lens: the market's worry might really be slower new sales or lower margins. We translate all of it into 'churn language'.", None),
]
for i, (t, f) in enumerate(lines, start=1):
    c = g.cell(i, 2, t)
    c.font = f or F_CALC
    c.alignment = Alignment(wrap_text=True, vertical="top")

# =============================================================== INPUTS
inp = wb.create_sheet("Inputs")
widths(inp, {"A": 46, **{c: 11 for c in "BCDEFGHI"}, "J": 70})
put(inp, "A1", "Inputs - Docebo Inc. (TSX: DCBO) reported figures", F(bold=True, size=14))
note(inp, "A2", "US$ millions unless stated. Blue = typed from filings. Source column says where each came from.")
header(inp, 4, ["Market & balance sheet", "Value", "Unit", "", "", "", "", "", "", "Source / note"])
market = [
    (5, "Share price (TSX: DCBO)", 33.22, USD2, "C$", "TMX Money / Investing.com, close 17-Sep-2026. UPDATE before each run."),
    (6, "Shares outstanding", 24.95, NUM2, "millions", "Post-SIB count per Docebo SIB filing (Jul/Aug 2026): ~24,947,594 common shares"),
    (7, "USD/CAD exchange rate", 1.41, '0.0000', "C$ per US$", "Mid/late Sep-2026 range 1.406-1.420 (Trading Economics). Match the price date."),
    (8, "Market capitalization", "=B5*B6", USDM, "C$ m", "Price x shares"),
    (9, "Cash & cash equivalents", 45.7, USDM, "US$ m", "Q2-2026 release, as at 30-Jun-2026"),
    (10, "Total borrowings (debt)", 88.0, USDM, "US$ m", "Q2-2026 release, as at 30-Jun-2026"),
    (11, "Net debt", "=B10-B9", USDM, "US$ m", "Debt - cash"),
    (12, "Gross margin (Q2-2026)", 0.794, PCT, "%", "Q2-2026 release: 79.4% of revenue"),
    (13, "Net revenue retention (FY2025, reported)", 0.99, PCT, "%", "Q4-2025 call: net dollar retention 99% incl. AWS"),
    (14, "Net revenue retention ex-AWS (FY2025)", 1.01, PCT, "%", "Q4-2025 call: 101% excluding the AWS (largest OEM customer) wind-down"),
]
for r, lab, v, fmt, unit, src in market:
    put(inp, f"A{r}", lab)
    isf = isinstance(v, str)
    put(inp, f"B{r}", v, F_CALC if isf else F_IN, fmt,
        FILL_KEY if r in (5, 13) else None)
    put(inp, f"C{r}", unit, F("595959", size=9))
    note(inp, f"J{r}", src)

header(inp, 16, ["Quarterly history (US$ m)", "Q3-24", "Q4-24", "Q1-25", "Q2-25", "Q3-25",
                 "Q4-25", "Q1-26", "Q2-26", "Source / note"])
Q = "BCDEFGHI"
hist = {
    17: ("ARR (annual recurring revenue)", [214.1, 219.7, 225.1, 233.0, 235.6, 238.1, 248.9, 255.1],
         "Quarterly releases. Q2-25 derived: Q2-26 ARR $255.1m was +9.5% YoY -> 255.1/1.095"),
    18: ("Total revenue", [55.4, 57.0, 57.3, 61.0, 62.0, 62.8, 65.6, 68.7],
         "Releases. Q1-25..Q4-25 derived as reported FCF / reported FCF margin (cross-checks to FY25 FCF $38.4m = 15.8%)"),
    19: ("Adjusted EBITDA", [8.7, 9.5, None, 9.2, 12.4, 13.3, 11.2, 11.2],
         "Releases. Q1-26 is an ESTIMATE (~17% margin per Q1-26 slides x revenue). Q1-25 not needed."),
    20: ("Free cash flow (cash from ops - capex)", [4.5, 10.1, 9.0, 11.4, 5.7, 12.3, 27.6, 3.1],
         "Releases (company-defined FCF). Q1-26 boosted by customer prepayments; Q2-26 low - use LTM to smooth."),
}
for r, (lab, vals, src) in hist.items():
    put(inp, f"A{r}", lab)
    for col, v in zip(Q, vals):
        if v is not None:
            put(inp, f"{col}{r}", v, F_IN, USDM)
    note(inp, f"J{r}", src)
put(inp, "A21", "FCF margin (%)")
put(inp, "A22", "Adj. EBITDA margin (%)")
put(inp, "A23", "ARR growth vs. same quarter last year (%)")
for col in Q:
    put(inp, f"{col}21", f"=IF({col}18>0,{col}20/{col}18,0)", fmt=PCT)
    put(inp, f"{col}22", f'=IF({col}19="","",{col}19/{col}18)', fmt=PCT)
for i, col in enumerate(Q[4:]):
    put(inp, f"{col}23", f"={col}17/{Q[i]}17-1", fmt=PCT)

header(inp, 25, ["Derived: last twelve months (LTM = Q3-25 to Q2-26)", "Value", "Unit", "", "", "", "", "", "", "How"])
derived = [
    (26, "LTM revenue", "=SUM(F18:I18)", USDM, "Sum of last 4 quarters"),
    (27, "LTM adjusted EBITDA", "=SUM(F19:I19)", USDM, ""),
    (28, "LTM free cash flow", "=SUM(F20:I20)", USDM, ""),
    (29, "LTM EBITDA margin", "=B27/B26", PCT, "Starting point for the margin ramp"),
    (30, "LTM FCF margin", "=B28/B26", PCT, ""),
    (31, "FCF conversion now (FCF / EBITDA)", "=B28/B27", PCT, "How much of EBITDA becomes cash today"),
    (32, "ARR now (30-Jun-2026)", "=I17", USDM, "Model starting point"),
    (33, "ARR one year ago (30-Jun-2025)", "=E17", USDM, ""),
    (34, "Net new ARR, last 12 months", "=B32-B33", USDM, "Ending - starting"),
    (35, "Gross new-customer ARR, last 12 months", "=B34+(1-B13)*B33", USDM,
     "Net new ARR + ARR lost to net churn. Uses REPORTED retention, so it doesn't move when you change the churn lever"),
    (36, "Revenue per $1 of average ARR", "=B26/AVERAGE(B32,B33)", '0.000', "Revenue also includes services, so it is a bit above ARR"),
    (37, "Enterprise value the market pays (US$ m)", "=B8/B7+B11", USDM, "Market cap in US$ + net debt"),
    (38, "EV / ARR", "=B37/B32", MULT, "What investors pay per $1 of ARR"),
    (39, "EV / LTM free cash flow", "=B37/B28", MULT, ""),
]
for r, lab, f, fmt, how in derived:
    put(inp, f"A{r}", lab)
    put(inp, f"B{r}", f, fmt=fmt)
    note(inp, f"J{r}", how)
inp.freeze_panes = "B5"

# =============================================================== ASSUMPTIONS
a = wb.create_sheet("Assumptions")
widths(a, {"A": 3, "B": 52, "C": 12, "D": 90})
put(a, "B1", "Assumptions - forecast levers", F(bold=True, size=14))
note(a, "B2", "Change a blue cell and every tab updates. C5 (churn) is the lever the reverse DCF solves for.")
header(a, 4, ["", "Lever", "Value", "Why this number"])
levers = [
    (5, "Gross ARR churn per year (% of starting ARR lost)", 0.09, True,
     "Starts equal to the churn implied by reported retention (C16). Goal Seek changes THIS cell."),
    (6, "Expansion per year (% of starting ARR from upsells)", 0.08, False,
     "Not disclosed. Assumption typical of mid-market SaaS. With NRR 99% it implies ~9% gross churn."),
    (7, "Growth in new-customer ARR, Year 1", 0.08, False,
     "Management cites strongest gross bookings since 2021 and ARR re-acceleration in 2026."),
    (8, "Growth in new-customer ARR, Year 10", 0.025, False, "Fades in a straight line to long-run economy growth."),
    (9, "Adj. EBITDA margin, long-term target", 0.28, True,
     "LTM is ~19%; Q3-26 guidance implies ~23%. Mature profitable SaaS often runs 25-35%."),
    (10, "FCF conversion (FCF / EBITDA), long-term", 0.85, False,
     "Today ~100% (low cash taxes, prepayments). Falls as cash taxes rise."),
    (11, "Years to reach long-term margin & conversion", 5, False, "Straight-line ramp from today's LTM figures."),
    (12, "Gross margin (for display)", "=Inputs!B12", False, "Linked from Inputs"),
    (13, "Discount rate (required return per year)", 0.10, True, "Starting point per project brief. See Sensitivity tab for 8-12%."),
    (14, "Terminal growth rate (forever after Year 10)", 0.025, False, "About long-run inflation + real growth. Must be below the discount rate."),
]
for r, lab, v, key, why in levers:
    put(a, f"B{r}", lab)
    isf = isinstance(v, str)
    put(a, f"C{r}", v, F_LINK if isf else F_IN, '0' if r == 11 else PCT, FILL_KEY if key else None)
    note(a, f"D{r}", why)
put(a, "B16", "Reported-equivalent churn = expansion + (1 - reported NRR)", bold=True)
put(a, "C16", "=C6+(1-Inputs!B13)", fmt=PCT, fill=FILL_ANS, bold=True)
note(a, "D16", "What the company's own retention number implies. The benchmark we compare the market's bet against.")
put(a, "B17", "Net revenue retention used in the DCF = 1 - churn + expansion")
put(a, "C17", "=1-C5+C6", fmt=PCT)

# =============================================================== ARR MODEL
m = wb.create_sheet("ARR Model")
widths(m, {"A": 46, "B": 13, **{c: 11 for c in YC}, "M": 3})
put(m, "A1", "ARR Model - 10-year forecast (US$ m)", F(bold=True, size=14))
for c in ["A"] + ["B"] + YC:
    m[f"{c}2"].fill = FILL_HDR
    m[f"{c}3"].fill = FILL_HDR
put(m, "A2", "Year", F_HDR)
put(m, "B2", 0, F_HDR, align="center")
put(m, "A3", "Period", F_HDR)
put(m, "B3", "Now (Jun-26)", F_HDR, align="center")
for i, c in enumerate(YC):
    prev = "B" if i == 0 else YC[i - 1]
    put(m, f"{c}2", f"={prev}2+1", F_HDR, align="center")
    put(m, f"{c}3", f'="Yr "&{c}2', F_HDR, align="center")
m.freeze_panes = "C4"
A = "Assumptions"


def mrow(r, lab, formula, fmt=USDM, font=F_CALC, bold=False, opening=None, border=None):
    put(m, f"A{r}", lab, bold=bold)
    for i, c in enumerate(YC):
        p = "B" if i == 0 else YC[i - 1]
        put(m, f"{c}{r}", formula.format(c=c, p=p), font, fmt, bold=bold, border=border)
    if opening is not None:
        put(m, f"B{r}", opening, F_LINK, fmt, bold=bold, border=border)
    if border:
        m[f"A{r}"].border = border


section(m, 5, "Drivers (ramp from today's numbers to long-term targets)", 12)
mrow(6, "Ramp progress (0% = today, 100% = long-term)", f"=MIN({{c}}$2/{A}!$C$11,1)", PCT)
mrow(7, "Growth in new-customer ARR", f"={A}!$C$7+({A}!$C$8-{A}!$C$7)*({{c}}$2-1)/9", PCT, F_LINK)
mrow(8, "Adj. EBITDA margin", f"=Inputs!$B$29+({A}!$C$9-Inputs!$B$29)*{{c}}6", PCT, F_LINK,
     opening="=Inputs!B29")
mrow(9, "FCF conversion (FCF / EBITDA)", f"=Inputs!$B$31+({A}!$C$10-Inputs!$B$31)*{{c}}6", PCT, F_LINK,
     opening="=Inputs!B31")
mrow(10, "FCF margin (EBITDA margin x conversion)", "={c}8*{c}9", PCT)

section(m, 12, "ARR bridge  (Starting ARR + new + expansion - churn = Ending ARR)", 12)
mrow(13, "Starting ARR", "={p}17")
mrow(14, "(+) New-customer ARR", "={p}14*(1+{c}7)", opening="=Inputs!B35")
mrow(15, "(+) Expansion ARR (upsells)", f"={{c}}13*{A}!$C$6", font=F_LINK)
mrow(16, "(-) Churned ARR", f"=-{{c}}13*{A}!$C$5", font=F_LINK)
mrow(17, "Ending ARR", "=SUM({c}13:{c}16)", bold=True, border=TOP, opening="=Inputs!B32")
mrow(18, "ARR growth (%)", "={c}17/{p}17-1", PCT)
mrow(19, "Net revenue retention (1 - churn + expansion)", "=({c}13+{c}15+{c}16)/{c}13", PCT)
m["B14"].comment = Comment("Year-0 value = gross new-customer ARR over the last 12 months (Inputs!B35).", "Model")

section(m, 21, "Revenue to free cash flow", 12)
mrow(22, "Revenue (avg ARR x revenue per $ of ARR)", "=AVERAGE({p}17,{c}17)*Inputs!$B$36", bold=True,
     opening="=Inputs!B26")
mrow(23, "Revenue growth (%)", "={c}22/{p}22-1", PCT)
mrow(24, "Gross profit", f"={{c}}22*{A}!$C$12")
mrow(25, "(-) Operating expenses (S&M, R&D, G&A)", "={c}26-{c}24")
mrow(26, "Adjusted EBITDA (operating profit before D&A)", "={c}22*{c}8", bold=True, border=TOP)
mrow(27, "(-) Cash taxes, capex & other", "=-{c}26*(1-{c}9)")
mrow(28, "Free cash flow", "={c}26+{c}27", bold=True, border=TOP, opening="=Inputs!B28")
m["B22"].comment = Comment("Year 0 column = last-twelve-months actuals.", "Model")

# =============================================================== DCF
d = wb.create_sheet("DCF")
widths(d, {"A": 46, "B": 13, **{c: 11 for c in YC}, "N": 80})
put(d, "A1", "DCF - what are those future cash flows worth today?", F(bold=True, size=14))
for c in ["A", "B"] + YC:
    d[f"{c}3"].fill = FILL_HDR
put(d, "A3", "Year", F_HDR)
for c in YC:
    put(d, f"{c}3", f"='ARR Model'!{c}2", F_HDR, align="center")
put(d, "A4", "Free cash flow (US$ m)")
put(d, "A5", "Discount factor = 1 / (1 + discount rate)^year")
put(d, "A6", "Present value of FCF (today's dollars)")
for c in YC:
    put(d, f"{c}4", f"='ARR Model'!{c}28", F_LINK, USDM)
    put(d, f"{c}5", f"=1/(1+{A}!$C$13)^{c}3", fmt='0.000')
    put(d, f"{c}6", f"={c}4*{c}5", fmt=USDM)
dcf = [
    (8, "Sum of PV of FCF, Years 1-10", f"=SUM({Y1}6:{Y10}6)", USDM, "Cash we can see"),
    (9, "Terminal value at Year 10", f"={Y10}4*(1+{A}!C14)/({A}!C13-{A}!C14)", USDM,
     "Value of every year after Year 10: FCF(Yr 11) / (discount rate - growth)"),
    (10, "PV of terminal value", f"=C9*{Y10}5", USDM, "Shrunk back to today's dollars"),
    (11, "Enterprise value (US$ m)", "=C8+C10", USDM, "Value of the business itself"),
    (12, "(+) Cash", "=Inputs!B9", USDM, ""),
    (13, "(-) Debt", "=-Inputs!B10", USDM, ""),
    (14, "Equity value (US$ m)", "=SUM(C11:C13)", USDM, "What belongs to shareholders"),
    (15, "Shares outstanding (m)", "=Inputs!B6", NUM2, ""),
    (16, "Value per share (US$)", "=C14/C15", USD2, ""),
    (17, "USD/CAD", "=Inputs!B7", '0.0000', ""),
    (18, "Model value per share (C$)", "=C16*C17", USD2, "Compare with the TSX price"),
    (19, "Market price today (C$)", "=Inputs!B5", USD2, ""),
    (20, "Upside / (downside) vs. market", "=C18/C19-1", PCT, "Positive = model says the stock is cheap"),
    (21, "Terminal value as % of enterprise value", "=C10/C11", PCT, "Usually 60-80% for a growing company"),
    (22, "GOAL SEEK TARGET -> model price (C$)", "=C18", USD2,
     "Goal Seek: set C22 to the market price by changing Assumptions!C5"),
]
for r, lab, f, fmt, why in dcf:
    bold = r in (11, 14, 18, 22)
    put(d, f"A{r}", lab, bold=bold)
    font = F_LINK if f.startswith("=Inputs") else F_CALC
    put(d, f"C{r}", f, font, fmt, bold=bold, fill=FILL_ANS if r in (18, 22) else None,
        border=TOP if r in (11, 14) else None)
    note(d, f"D{r}", why)

# =============================================================== REVERSE DCF
rv = wb.create_sheet("Reverse DCF")
widths(rv, {"A": 12, **{c: 9 for c in "BCDEFGHIJKL"}, "M": 11, "N": 11})
rv.column_dimensions["A"].width = 12
put(rv, "A1", "Reverse DCF - the churn rate today's price is betting on", F(bold=True, size=14))
note(rv, "A2", "Solved with a live formula: the price is computed for churn from 0% to 50% below, then we find where it crosses the market price.")

LAD0, STEP, NSTEP = 30, 0.005, 101     # ladder rows 30..130, churn 0%..50%
LAD1 = LAD0 + NSTEP - 1
res = [
    (4, "Market price today (C$)", "=Inputs!B5", USD2),
    (5, "MARKET-IMPLIED CHURN (per year)", None, PCT),
    (6, "Reported-equivalent churn (per year)", f"={A}!C16", PCT),
    (7, "Gap: implied minus reported (percentage points)", "=IFERROR(H5-H6,\"n/a\")", '+0.0%;-0.0%;0.0%'),
    (8, "Market-implied net revenue retention", f"=IFERROR(1-H5+{A}!C6,\"n/a\")", PCT),
    (9, "Reported net revenue retention (FY2025)", "=Inputs!B13", PCT),
    (10, "Model price at reported churn (C$)", f"=N{LAD1 + 3}", USD2),
    (11, "Check: model price at implied churn (C$) - should equal market", f"=N{LAD1 + 4}", USD2),
]
for r, lab, f, fmt in res:
    rv.merge_cells(f"A{r}:G{r}")
    put(rv, f"A{r}", lab, bold=r in (5, 7))
    if f:
        put(rv, f"H{r}", f, F_LINK if "!" in f and not f.startswith("=IF") else F_CALC, fmt,
            bold=r in (5, 7))
# implied churn via interpolation on the descending price ladder
px = f"$N${LAD0}:$N${LAD1}"
ch = f"$A${LAD0}:$A${LAD1}"
k = f"MATCH(H4,{px},-1)"
implied = (f'=IF(H4>N{LAD0},"below 0%",IF(H4<N{LAD1},"above 50%",'
           f"INDEX({ch},{k})+(INDEX({px},{k})-H4)/(INDEX({px},{k})-INDEX({px},{k}+1))*{STEP}))")
put(rv, "H5", implied, fmt=PCT, fill=FILL_ANS, bold=True, border=BOX)
for r in (6, 7, 8, 9):
    rv[f"H{r}"].fill = FILL_ANS
rv.merge_cells("A13:N13")
verdict = ('=IF(ISNUMBER(H5),IF(H5>H6,"PESSIMISTIC: the price only makes sense if Docebo loses about "&TEXT(H5,"0%")'
           '&" of its ARR every year - vs. about "&TEXT(H6,"0%")&" implied by what it reports. The market is betting things get worse (or doubts the numbers).",'
           '"OPTIMISTIC: the price needs churn of only "&TEXT(H5,"0.0%")&" - better than the "&TEXT(H6,"0.0%")&" the company reports. The market expects improvement."),'
           '"Implied churn is outside 0-50%: check the inputs.")')
put(rv, "A13", verdict, F(bold=True, size=11, color="C00000"))
rv["A13"].alignment = Alignment(wrap_text=True, vertical="top")
rv.row_dimensions[13].height = 48
note(rv, "A15", "How to read it: implied churn HIGHER than reported = market pessimistic (stock may be cheap if you trust the numbers). LOWER = market optimistic.")
note(rv, "A16", "Cross-check with Goal Seek (see Start Here). The two answers should agree to about 0.1 percentage point.")

# ladder engine
put(rv, f"A{LAD0 - 3}", "Engine: price for each churn rate (same maths as ARR Model + DCF, one row per churn)", bold=True)
header(rv, LAD0 - 1, ["Churn"] + [f"ARR Y{i}" for i in range(11)] + ["EV (US$m)", "Price C$"])
ARRC = [chr(ord("B") + i) for i in range(11)]   # B..L


def engine_row(ws, r, churn_formula, rate="Assumptions!$C$13", df_row=None):
    """ARR Y0..Y10 in B..L, EV in M, price C$ in N for a given churn formula."""
    put(ws, f"A{r}", churn_formula, F_IN if not str(churn_formula).startswith("=") else F_CALC, PCT)
    put(ws, f"B{r}", "='ARR Model'!$B$17", F_LINK, '#,##0')
    for i in range(1, 11):
        c, p, ym = ARRC[i], ARRC[i - 1], YC[i - 1]
        put(ws, f"{c}{r}", f"={p}{r}*(1-$A{r}+{A}!$C$6)+'ARR Model'!{ym}$14", fmt='#,##0')
    rev = f"(B{r}:K{r}+C{r}:L{r})/2*Inputs!$B$36"
    ev = (f"=SUMPRODUCT({rev},'ARR Model'!$C$10:$L$10,DCF!$C$5:$L$5)"
          f"+AVERAGE(K{r},L{r})*Inputs!$B$36*'ARR Model'!$L$10*(1+{A}!$C$14)/({A}!$C$13-{A}!$C$14)*DCF!$L$5")
    put(ws, f"M{r}", ev, fmt='#,##0')
    put(ws, f"N{r}", f"=(M{r}+Inputs!$B$9-Inputs!$B$10)/Inputs!$B$6*Inputs!$B$7", fmt=USD2)


for i in range(NSTEP):
    r = LAD0 + i
    engine_row(rv, r, round(i * STEP, 4))
    rv[f"A{r}"].font = F_CALC
put(rv, f"A{LAD1 + 2}", "Checks", bold=True)
engine_row(rv, LAD1 + 3, f"={A}!C16")
engine_row(rv, LAD1 + 4, "=IF(ISNUMBER(H5),H5,0)")
put(rv, f"O{LAD1 + 3}", "<- at reported churn", F("595959", italic=True, size=9))
put(rv, f"O{LAD1 + 4}", "<- at implied churn", F("595959", italic=True, size=9))
engine_row(rv, LAD1 + 5, f"={A}!C5")
put(rv, f"O{LAD1 + 5}", "<- at the churn in Assumptions!C5 (must equal DCF!C18)", F("595959", italic=True, size=9))
put(rv, f"A{LAD1 + 6}", f'=IF(ABS(N{LAD1 + 5}-DCF!C18)<0.005,"CHECK OK: engine matches DCF tab","CHECK FAILED: engine differs from DCF tab")', bold=True)
rv.freeze_panes = "A4"

# =============================================================== SENSITIVITY
s = wb.create_sheet("Sensitivity")
widths(s, {"A": 30, **{c: 12 for c in "BCDEFG"}})
put(s, "A1", "Sensitivity - model share price (C$) for each churn x discount rate", F(bold=True, size=14))
note(s, "A2", "Green = model price at or above today's market price (stock looks cheap). Red = below (stock looks expensive). Blue headers are editable.")
put(s, "A3", "Market price (C$):")
put(s, "B3", "=Inputs!B5", F_LINK, USD2, bold=True)
rates = [0.08, 0.09, 0.10, 0.11, 0.12]
churns = [0.05, 0.075, 0.09, 0.10, 0.125, 0.15, 0.20, 0.25, 0.30]
put(s, "A5", "Churn per year ↓  /  Discount rate →", F_HDR, fill=FILL_HDR)
GC = "BCDEF"
for j, rt in enumerate(rates):
    put(s, f"{GC[j]}5", rt, F("FFFFFF", True), PCT, FILL_HDR, align="center")
E0 = 20  # engine rows
for i, chv in enumerate(churns):
    r = 6 + i
    put(s, f"A{r}", chv, F_IN, PCT, align="center")
    er = E0 + 2 + i
    for j in range(len(rates)):
        rt = f"{GC[j]}$5"
        f = (f"=(SUMPRODUCT($M{er}:$V{er},'ARR Model'!$C$10:$L$10,1/(1+{rt})^'ARR Model'!$C$2:$L$2)"
             f"+$V{er}*'ARR Model'!$L$10*(1+{A}!$C$14)/({rt}-{A}!$C$14)/(1+{rt})^10"
             f"+Inputs!$B$9-Inputs!$B$10)/Inputs!$B$6*Inputs!$B$7")
        put(s, f"{GC[j]}{r}", f, fmt=USD2, align="center")
last = 5 + len(churns)
rng = f"B6:F{last}"
s.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["$B$3"],
                                             fill=PatternFill("solid", fgColor="C6EFCE")))
s.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["$B$3"],
                                             fill=PatternFill("solid", fgColor="FFC7CE")))
note(s, f"A{last + 2}", "Read across a row: a higher discount rate (investors demand more return) lowers the price.")
note(s, f"A{last + 3}", "Read down a column: more churn lowers the price. Find the cells closest to today's price - that's the market's bet.")
note(s, f"A{last + 4}", "Reported-equivalent churn is ~9%; market-implied churn at 10% is on the Reverse DCF tab.")
# engine: ARR Y0..Y10 in B..L, revenue Y1..Y10 in M..V
put(s, f"A{E0}", "Engine (one row per churn rate): ARR Y0-Y10, then revenue Y1-Y10", bold=True)
header(s, E0 + 1, ["Churn"] + [f"ARR Y{i}" for i in range(11)] + [f"Rev Y{i}" for i in range(1, 11)])
RC = [chr(ord("M") + i) for i in range(10)]  # M..V
for i in range(len(churns)):
    er = E0 + 2 + i
    put(s, f"A{er}", f"=A{6 + i}", fmt=PCT)
    put(s, f"B{er}", "='ARR Model'!$B$17", F_LINK, '#,##0')
    for t in range(1, 11):
        c, p = ARRC[t], ARRC[t - 1]
        put(s, f"{c}{er}", f"={p}{er}*(1-$A{er}+{A}!$C$6)+'ARR Model'!{YC[t - 1]}$14", fmt='#,##0')
        put(s, f"{RC[t - 1]}{er}", f"=AVERAGE({p}{er},{c}{er})*Inputs!$B$36", fmt='#,##0')
for c in "GHIJKLMNOPQRSTUV":
    s.column_dimensions[c].width = 9

# =============================================================== TRACKER
t = wb.create_sheet("Tracker")
widths(t, {"A": 13, "B": 40, "C": 11, "D": 10, "E": 11, "F": 10, "G": 13, "H": 13, "I": 60})
put(t, "A1", "Tracker - did the market's hidden bet change after news?", F(bold=True, size=14))
note(t, "A2", "After each earnings release: update Inputs, then add a row and PASTE AS VALUES the live results from Reverse DCF (Ctrl+Shift+V).")
header(t, 4, ["Date", "Event", "Price C$", "USD/CAD", "ARR US$m", "NRR", "Implied churn", "Implied NRR", "What changed / my read"])
put(t, "A5", "2026-09-17", F_IN)
put(t, "B5", "Baseline (after Q2-2026 results)", F_IN)
for col, f, fmt in (("C", "=Inputs!B5", USD2), ("D", "=Inputs!B7", '0.00'), ("E", "=Inputs!B32", USDM),
                    ("F", "=Inputs!B13", PCT), ("G", "='Reverse DCF'!H5", PCT),
                    ("H", "='Reverse DCF'!H8", PCT)):
    put(t, f"{col}5", f, F_LINK, fmt)
put(t, "I5", "Live link for now - paste as values before adding the next row.", F("595959", italic=True, size=9))
put(t, "A6", "Nov-2026", F_IN)
put(t, "B6", "Q3-2026 results (expected early Nov 2026)", F_IN)
put(t, "I6", "Fill in after the release.", F("595959", italic=True, size=9))

wb.move_sheet("Reverse DCF", offset=-3)   # put the answer right after Assumptions
wb.calculation.fullCalcOnLoad = True
wb.save("Docebo_Reverse_DCF.xlsx")
print("Saved Docebo_Reverse_DCF.xlsx")
